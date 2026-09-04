"""
Export utilities — turns stored review history into downloadable
Excel (.xlsx) or PDF reports for the dashboard's export buttons.
"""

import io
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer


COLUMNS = ["ID", "Review", "Rating", "Label", "Confidence", "Level", "Human Review?", "Reasoning", "Date"]


def _row_values(r: dict) -> list:
    return [
        r["id"],
        r["review_text"],
        r.get("rating") or "-",
        r["predicted_label"],
        f'{r["confidence"]}%',
        r.get("confidence_level") or "-",
        "Yes" if r.get("needs_human_review") else "No",
        r.get("reasoning") or "",
        str(r["created_at"]),
    ]


def build_excel(rows: list[dict]) -> bytes:
    """Builds an .xlsx workbook from review history rows and returns raw bytes."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Review Predictions"

    header_fill = PatternFill(start_color="2E5090", end_color="2E5090", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)

    ws.append(COLUMNS)
    for col_idx in range(1, len(COLUMNS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")

    fake_fill = PatternFill(start_color="FDE8E8", end_color="FDE8E8", fill_type="solid")
    genuine_fill = PatternFill(start_color="E8F6EE", end_color="E8F6EE", fill_type="solid")

    for r in rows:
        ws.append(_row_values(r))
        label_cell = ws.cell(row=ws.max_row, column=4)
        label_cell.fill = fake_fill if r["predicted_label"] == "Fake" else genuine_fill

    widths = [6, 45, 8, 10, 12, 9, 13, 40, 20]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def build_pdf(rows: list[dict], title: str = "Fake Review Detection — History Report") -> bytes:
    """Builds a landscape PDF table report from review history rows and returns raw bytes."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=1.5 * cm, rightMargin=1.5 * cm,
        topMargin=1.5 * cm, bottomMargin=1.5 * cm,
    )
    styles = getSampleStyleSheet()
    elements = [
        Paragraph(title, styles["Title"]),
        Paragraph(f"Generated on {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"]),
        Spacer(1, 12),
    ]

    def truncate(text, n=60):
        text = str(text)
        return text if len(text) <= n else text[:n] + "…"

    data = [COLUMNS]
    for r in rows:
        row = _row_values(r)
        row[1] = truncate(row[1], 45)   # Review
        row[7] = truncate(row[7], 40)   # Reasoning
        data.append(row)

    col_widths = [1 * cm, 6.5 * cm, 1.3 * cm, 1.8 * cm, 1.8 * cm, 1.6 * cm, 2 * cm, 6 * cm, 3 * cm]
    table = Table(data, colWidths=col_widths, repeatRows=1)

    style_commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#2E5090")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#CCCCCC")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F7F7F7")]),
    ]
    for i, r in enumerate(rows, start=1):
        if r["predicted_label"] == "Fake":
            style_commands.append(("TEXTCOLOR", (3, i), (3, i), colors.HexColor("#C0392B")))
        else:
            style_commands.append(("TEXTCOLOR", (3, i), (3, i), colors.HexColor("#1B7A3D")))

    table.setStyle(TableStyle(style_commands))
    elements.append(table)

    doc.build(elements)
    buffer.seek(0)
    return buffer.getvalue()
